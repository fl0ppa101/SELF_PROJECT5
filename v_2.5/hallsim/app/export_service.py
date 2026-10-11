from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PyQt6.QtCore import QPoint, QRectF, Qt
from PyQt6.QtGui import QColor, QFont, QImage, QPainter

from hallsim import __version__
from hallsim.core.enums import MeasurementSource, SensorStudyAxis
from hallsim.visualization.theme import get_theme


@dataclass
class ExportData:
    title: str
    headers: list
    rows: list
    parameters: list


def scene_parameters(scene, catalog, *, sensor=True, path=True):
    rows = [
        ("Дата экспорта", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Версия HallSim", __version__),
        ("Модель", "Точечные магнитные диполи; идеальный эффект Холла"),
    ]
    for magnet in scene.magnets:
        definition = catalog[magnet.definition_id]
        name = f"Магнит {magnet.id}"
        rows.extend(
            [
                (f"{name}: тип", definition.display_name),
                (f"{name}: момент, А·м²", definition.moment_Am2),
                (f"{name}: включён", "Да" if magnet.enabled else "Нет"),
                (f"{name}: X, мм", magnet.x_mm),
                (f"{name}: Y, мм", magnet.y_mm),
                (f"{name}: угол S→N, °", magnet.angle_deg),
            ]
        )
    if path:
        p = scene.path
        rows.extend(
            [
                ("A: X, мм", p.ax_mm),
                ("A: Y, мм", p.ay_mm),
                ("B: X, мм", p.bx_mm),
                ("B: Y, мм", p.by_mm),
            ]
        )
    if sensor:
        s = scene.sensor
        rows.extend(
            [
                ("Датчик: X, мм", s.x_mm),
                ("Датчик: Y, мм", s.y_mm),
                ("Датчик: угол чувствительной оси, °", s.angle_deg),
            ]
        )
    return rows


def hall_parameters(params):
    return [
        ("Ток I, мА", params.current_mA),
        ("Толщина d, мм", params.thickness_mm),
        ("Коэффициент Холла R_H, м³/Кл", params.hall_coefficient_m3_C),
    ]


def profile_data(scene, profile, catalog):
    if profile is None or not scene.path.enabled:
        raise ValueError("Дождитесь расчёта профиля для текущего пути")
    p = scene.path
    length = math.hypot(p.bx_mm - p.ax_mm, p.by_mm - p.ay_mm)
    rows = []
    for distance, value, valid in zip(
        profile.distance_mm, profile.modB_mT, profile.valid_mask
    ):
        t = float(distance) / length
        rows.append(
            (
                float(distance),
                p.ax_mm + t * (p.bx_mm - p.ax_mm),
                p.ay_mm + t * (p.by_mm - p.ay_mm),
                float(value) if valid else None,
            )
        )
    parameters = scene_parameters(scene, catalog)
    parameters.append(("Количество точек профиля", len(rows)))
    return ExportData(
        "Профиль магнитного поля A–B",
        ["s, мм", "X, мм", "Y, мм", "|B|, мТл"],
        rows,
        parameters,
    )


def experiment_data(controller):
    if controller.running or not controller.state.points:
        raise ValueError("Сначала завершите измерения")
    headers = [
        "№",
        "s, мм",
        "X, мм",
        "Y, мм",
        "|U_H|max, мВ",
        "|B|изм, мТл",
        "Источник",
        "Число замеров",
    ]
    comparison = controller.comparison
    if comparison is not None:
        headers.extend(["|B|теор, мТл", "Абсолютное отклонение, мТл"])
    rows = []
    for index, point in enumerate(
        sorted(controller.state.points, key=lambda p: p.s_mm)
    ):
        source = (
            "Ручной" if point.source is MeasurementSource.MANUAL else "Автоматический"
        )
        row = [
            index + 1,
            point.s_mm,
            point.x_mm,
            point.y_mm,
            point.max_abs_hall_voltage_mV,
            point.measured_modB_mT,
            source,
            point.measurement_count,
        ]
        if comparison is not None:
            row.extend(
                [
                    float(comparison.theoretical_at_measured_mT[index]),
                    float(comparison.abs_error_mT[index]),
                ]
            )
        rows.append(row)
    parameters = scene_parameters(
        controller.snapshot,
        controller.scene_controller.magnet_definitions(),
        sensor=False,
    )
    parameters.extend(hall_parameters(controller.state.sensor_params))
    parameters.extend(
        [
            (
                "Текущее число точек автоматического прохода",
                controller.state.config.automatic_point_count,
            ),
            (
                "Текущий шаг автоматического поворота, °",
                controller.state.config.angular_step_deg,
            ),
            ("Результаты", "Завершённые точки; могут объединять несколько проходов"),
        ]
    )
    if comparison is not None:
        parameters.append(
            ("Среднее абсолютное отклонение, мТл", comparison.mean_abs_error_mT)
        )
    return ExportData(
        "Эксперимент: измерение магнитного поля", headers, rows, parameters
    )


def study_data(controller):
    controller.timer.stop()
    controller._compute_curve()
    curve = controller.curve
    if curve is None:
        raise ValueError("Сначала постройте график исследования")
    names = {
        SensorStudyAxis.CURRENT: "I",
        SensorStudyAxis.THICKNESS: "d",
        SensorStudyAxis.ANGLE: "α",
    }
    axis = names[curve.axis]
    rows = [(float(x), float(y)) for x, y in zip(curve.x_values, curve.hall_voltage_mV)]
    parameters = scene_parameters(
        controller.scene(), controller.scene_controller.magnet_definitions(), path=False
    )
    parameters.extend(hall_parameters(controller.state.sensor_params))
    parameters.extend(
        [
            ("Изменяемый параметр", axis),
            (f"Начало диапазона, {curve.x_unit}", float(curve.x_values[0])),
            (f"Конец диапазона, {curve.x_unit}", float(curve.x_values[-1])),
            ("Количество точек", len(rows)),
        ]
    )
    return ExportData(
        f"Исследование датчика: U_H({axis})",
        [f"{axis}, {curve.x_unit}", "U_H, мВ"],
        rows,
        parameters,
    )


def atomic_write(destination, writer):
    destination = Path(destination)
    with NamedTemporaryFile(
        dir=destination.parent,
        prefix=".hallsim-",
        suffix=destination.suffix,
        delete=False,
    ) as handle:
        temporary = Path(handle.name)
    try:
        writer(temporary)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def save_workbook(destination, data):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = data.title
    for name, headers, rows in (
        ("Результаты", data.headers, data.rows),
        (
            "Параметры",
            ["Параметр", "Значение"],
            [("Раздел", data.title), *data.parameters],
        ),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(
                [
                    (
                        "Недопустимое значение"
                        if value is None
                        or isinstance(value, float)
                        and not math.isfinite(value)
                        else value
                    )
                    for value in row
                ]
            )
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, float):
                    magnitude = abs(cell.value)
                    cell.number_format = (
                        "0.##############E+00"
                        if magnitude and (magnitude < 1e-6 or magnitude >= 1e9)
                        else "0.##############"
                    )
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="71372E")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.columns, 1):
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(
                65, max(16, width)
            )
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    atomic_write(destination, workbook.save)
    workbook.close()


def save_plot_image(destination, widget, title, theme):
    palette = get_theme(theme)
    selection = None
    hidden_items = []
    if hasattr(widget, "objects"):
        selection = widget.objects.selection
        widget.objects.set_selection(None)
        hidden_items.append(widget.experiment.preview)
        widget.path.set_path(widget.scene.path if widget.scene else None)
    for name in ("tooltip", "cross_x", "cross_y"):
        item = getattr(widget, name, None)
        if item is not None:
            hidden_items.append(item)
    visibility = [(item, item.isVisible()) for item in hidden_items]
    for item, _ in visibility:
        item.hide()
    try:
        widget.grab()
        width, height = widget.width(), widget.height()
        image = QImage(width * 2, (height + 48) * 2, QImage.Format.Format_ARGB32)
        image.fill(QColor(palette.background))
        painter = QPainter(image)
        try:
            painter.scale(2, 2)
            painter.setPen(QColor(palette.foreground))
            painter.setFont(QFont("Sans Serif", 12))
            painter.drawText(
                QRectF(12, 0, width - 24, 48), Qt.AlignmentFlag.AlignVCenter, title
            )
            painter.translate(0, 48)
            widget.render(painter, QPoint())
        finally:
            painter.end()

        def write(path):
            if not image.save(str(path), "PNG"):
                raise OSError("Не удалось сохранить изображение")

        atomic_write(destination, write)
    finally:
        for item, visible in visibility:
            item.setVisible(visible)
        if hasattr(widget, "objects"):
            widget.objects.set_selection(selection)
