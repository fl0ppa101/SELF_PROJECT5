import json
import math
import os
from pathlib import Path
import tempfile
import time
from threading import Event
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
import numpy as np
from PyQt6.QtCore import QSettings, Qt
from PyQt6.QtGui import QImage
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication, QMessageBox, QFileDialog

from hallsim.app.calculation_service import CalculationService
from hallsim.app.controller import AppController
from hallsim.app.export_service import (
    ExportData,
    atomic_write,
    experiment_data,
    save_workbook,
    study_data,
)
from hallsim.app.lab_controller import ExperimentController, SensorStudyController
from hallsim.app.lab_window import MainWindow
from hallsim.app.magnet_catalog_service import MagnetCatalogService
from hallsim.app.presets import single_magnet
from hallsim.app.settings_service import SettingsService
from hallsim.core.enums import SensorStudyAxis
from hallsim.core.contracts import FieldGrid
from hallsim.physics.engine import PhysicsEngine
from hallsim.visualization.streamline_layer import GridSampler, StreamlineBuilder


APP = QApplication.instance() or QApplication([])


def wait_until(predicate, seconds=5):
    deadline = time.monotonic() + seconds
    while not predicate() and time.monotonic() < deadline:
        QTest.qWait(10)
    if not predicate():
        raise AssertionError("Операция не завершилась вовремя")


class ReleaseTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="hallsim-test-")
        self.directory = Path(self.temporary.name)
        self.settings = SettingsService(
            QSettings(str(self.directory / "settings.ini"), QSettings.Format.IniFormat),
            self.directory / "scene.json",
        )
        self.catalog = MagnetCatalogService(self.directory / "magnets.json")
        self.engine = PhysicsEngine(self.catalog.definitions())
        self.visual = AppController(
            calculation_service=CalculationService(self.catalog.definitions()),
            catalog_service=self.catalog,
        )
        self.experiment = ExperimentController(self.engine, self.catalog)
        self.study = SensorStudyController(self.engine, self.catalog)
        self.window = MainWindow(
            self.visual, self.experiment, self.study, self.settings
        )
        self.window.resize(1366, 700)
        self.window.show()
        QTest.qWait(20)

    def tearDown(self):
        self.window.shutdown()
        self.window.hide()
        self.window.deleteLater()
        APP.processEvents()
        self.temporary.cleanup()

    def click_world(self, canvas, x, y):
        position = (canvas.world_to_screen(x, y) - canvas._viewport_origin()).toPoint()
        QTest.mouseClick(
            canvas.graphics.viewport(), Qt.MouseButton.LeftButton, pos=position
        )

    def measure(self, count=5):
        self.experiment.edit_scene("set_path", -40, 20, 40, 20)
        self.experiment.set_config(count, 10)
        self.experiment.start_automatic()
        wait_until(lambda: not self.experiment.running)

    def test_settings_visible_and_small_window(self):
        self.window.show_settings()
        QTest.qWait(20)
        self.assertTrue(self.window.visualization.settings_page.isVisible())
        self.window.settings_dialog.hide()
        for index in range(3):
            self.window.set_section(index)
            QTest.qWait(20)
            self.assertLessEqual(self.window.height(), 700)
        self.window.visualization.path_controls.toggle.setChecked(True)
        self.window.experiment.path_controls.toggle.setChecked(True)
        self.assertLessEqual(self.window.minimumSizeHint().height(), 700)

    def test_path_two_clicks_commit_once_and_cancel_preserves_scene(self):
        self.window.set_section(1)
        page = self.window.experiment
        before = self.experiment.scene().path
        page.path_button.setChecked(True)
        self.click_world(page.canvas, -40, 20)
        self.assertEqual(self.experiment.scene().path, before)
        self.assertEqual(page.canvas._path_stage, "B")
        self.assertFalse(page.auto_button.isEnabled())
        page.canvas.cancel_action()
        self.assertEqual(self.experiment.scene().path, before)
        self.assertFalse(page.path_button.isChecked())
        page.path_button.setChecked(True)
        self.click_world(page.canvas, -40, 20)
        self.click_world(page.canvas, 40, 20)
        path = self.experiment.scene().path
        self.assertTrue(path.enabled)
        self.assertAlmostEqual(path.ax_mm, -40, delta=1)
        self.assertAlmostEqual(path.bx_mm, 40, delta=1)
        self.assertEqual(self.experiment.scene().revision, 1)
        self.assertIsNone(page.canvas._path_stage)

    def test_escape_shortcut_cancels_before_and_after_first_point(self):
        self.window.activateWindow()
        self.window.set_section(1)
        QTest.qWait(30)
        page = self.window.experiment
        before = self.experiment.scene().path
        for first_click in (False, True):
            page.path_button.setChecked(True)
            if first_click:
                self.click_world(page.canvas, -40, 20)
            QTest.keyClick(page.canvas, Qt.Key.Key_Escape)
            QTest.qWait(10)
            self.assertIsNone(page.canvas._path_stage)
            self.assertFalse(page.path_button.isChecked())
            self.assertEqual(self.experiment.scene().path, before)

    def test_export_dialog_save_cancel_and_failure(self):
        destination = self.directory / "dialog.xlsx"
        button = self.window.study.export_button

        def accept(dialog):
            self.assertFalse(dialog.testOption(QFileDialog.Option.DontConfirmOverwrite))
            self.assertEqual(dialog.defaultSuffix(), "xlsx")
            dialog.selectFile(str(destination))
            return QFileDialog.DialogCode.Accepted

        with patch.object(QFileDialog, "exec", accept), patch.object(
            QMessageBox, "information"
        ) as info:
            button._export("xlsx", "sensor", self.window.study._prepare_table_export)
            info.assert_called_once()
        self.assertTrue(destination.exists())
        before = destination.read_bytes()
        with patch.object(
            QFileDialog, "exec", return_value=QFileDialog.DialogCode.Rejected
        ):
            button._export("xlsx", "sensor", self.window.study._prepare_table_export)
        self.assertEqual(destination.read_bytes(), before)

        def failure():
            raise OSError("Нет доступа к файлу")

        with patch.object(QMessageBox, "warning") as warning:
            button._export("xlsx", "sensor", failure)
            warning.assert_called_once()

    def test_numeric_path_clear_and_reject_identical_points(self):
        controls = self.window.visualization.path_controls
        for control, value in zip(controls.coordinates, (-30, 15, 50, 25)):
            control.setValue(value)
        controls.apply_button.click()
        self.assertEqual(self.visual.scene().path.ax_mm, -30)
        previous = self.visual.scene().path
        with self.assertRaises(ValueError):
            self.visual.set_path(1, 2, 1, 2)
        self.assertEqual(self.visual.scene().path, previous)
        controls.clear_button.click()
        self.assertFalse(self.visual.scene().path.enabled)
        self.assertIsNone(self.window.visualization.path_plot.profile)

    def test_old_profile_cleared_when_calculation_fails(self):
        self.visual.set_path(-30, 20, 30, 20)
        plot = self.window.visualization.path_plot
        wait_until(lambda: plot.profile is not None)
        self.visual.set_path_point("B", -30, 20)
        self.assertIsNone(plot.profile)
        wait_until(
            lambda: "различаться"
            in self.window.visualization.statusBar().currentMessage()
        )
        self.assertIsNone(plot.profile)
        self.assertIsNone(plot.curve.xData)

    def test_field_layers_are_independent_and_follow_zoom(self):
        canvas = self.window.visualization.canvas
        wait_until(lambda: canvas.grid is not None)
        self.assertTrue(canvas.vectors.item.isVisible())
        self.assertFalse(canvas.streamlines.item.isVisible())
        self.assertFalse(canvas.vectors.item.path().isEmpty())
        self.visual.set_layer_visibility("field_lines", True)
        wait_until(lambda: bool(canvas.streamlines.geometry))
        self.assertFalse(canvas.streamlines.item.path().isEmpty())
        self.assertTrue(canvas.streamlines.item.isVisible())
        self.assertTrue(canvas.vectors.item.isVisible())
        self.visual.set_layer_visibility("field_arrows", False)
        self.assertFalse(canvas.vectors.item.isVisible())
        self.visual.set_layer_visibility("field_arrows", True)
        before = canvas.vectors.item.path().boundingRect()
        canvas.zoom_in()
        after = canvas.vectors.item.path().boundingRect()
        self.assertNotEqual(before, after)
        self.visual.set_layer_visibility("heatmap", False)
        self.assertNotEqual(canvas.vectors.item.pen().color().name(), "#e4f5ef")
        self.visual.set_layer_visibility("field_lines", False)

    def test_mouse_drag_keeps_preview_and_finishes_with_full_grid(self):
        page = self.window.visualization
        canvas = page.canvas
        wait_until(lambda: canvas.grid is not None)
        original = canvas.grid
        viewport = canvas.graphics.viewport()

        def position(x, y):
            return (canvas.world_to_screen(x, y) - canvas._viewport_origin()).toPoint()

        QTest.mousePress(viewport, Qt.MouseButton.LeftButton, pos=position(0, 0))
        try:
            QTest.mouseMove(viewport, position(15, 10))
            self.assertIs(canvas.grid, original)
            self.assertFalse(canvas.empty_label.isVisible())
            page._update_exports()
            self.assertFalse(page.export_map.isEnabled())
            with self.assertRaises(ValueError):
                page._prepare_field_image()
            wait_until(lambda: canvas.grid.bx_mT.shape == (200, 320))
            self.assertFalse(canvas.empty_label.isVisible())
            self.assertFalse(page._field_image_ready())
        finally:
            QTest.mouseRelease(
                viewport, Qt.MouseButton.LeftButton, pos=position(15, 10)
            )
        wait_until(page._field_image_ready)
        self.assertEqual(canvas.grid.bx_mT.shape, (400, 640))
        self.assertEqual(canvas.grid.scene_revision, self.visual.scene().revision)
        self.assertAlmostEqual(self.visual.scene().magnets[0].x_mm, 15, delta=1)
        self.assertFalse(canvas.empty_label.isVisible())

    def test_disabling_lines_cancels_build_and_does_not_restart_on_release(self):
        page = self.window.visualization
        canvas = page.canvas
        wait_until(lambda: canvas.grid is not None)
        started, release = Event(), Event()
        cancellations = []
        build = StreamlineBuilder().build

        def delayed(grid, **kwargs):
            cancellations.append(kwargs["cancelled"])
            started.set()
            release.wait(5)
            return build(grid, **kwargs)

        with patch.object(StreamlineBuilder, "build", side_effect=delayed) as builder:
            try:
                self.visual.set_layer_visibility("field_lines", True)
                wait_until(started.is_set)
                page._update_exports()
                self.assertFalse(page.export_map.isEnabled())
                self.visual.set_layer_visibility("field_arrows", False)
                self.visual.set_layer_visibility("field_arrows", True)
                self.assertFalse(cancellations[0].is_set())
                self.visual.set_layer_visibility("field_lines", False)
                self.assertTrue(cancellations[0].is_set())
                canvas.set_interaction_quality("drag")
                canvas.set_interaction_quality("full")
            finally:
                release.set()
            wait_until(lambda: canvas.streamlines.running is None)
            QTest.qWait(30)
            self.assertEqual(builder.call_count, 1)
        self.assertEqual(canvas.streamlines.geometry, [])
        self.assertTrue(canvas.streamlines.item.path().isEmpty())
        self.assertTrue(page._field_image_ready())

    def test_profile_export_after_sensor_edit(self):
        self.visual.set_path(-40, 20, 40, 20)
        wait_until(lambda: self.window.visualization.path_plot.profile is not None)
        self.visual.set_sensor_position(10, 30)
        destination = self.directory / "profile.xlsx"
        self.window.visualization._prepare_profile_export()(destination)
        book = load_workbook(destination)
        self.assertEqual(book.sheetnames, ["Результаты", "Параметры"])
        self.assertEqual(book["Результаты"].max_row, 301)
        self.assertEqual(book["Результаты"].cell(2, 2).value, -40)
        self.assertEqual(book["Результаты"].cell(301, 2).value, 40)
        book.close()

    def test_experiment_export_theory_only_after_comparison(self):
        self.measure()
        self.assertEqual(len(experiment_data(self.experiment).headers), 8)
        self.experiment.compare()
        destination = self.directory / "comparison.xlsx"
        save_workbook(destination, experiment_data(self.experiment))
        book = load_workbook(destination)
        sheet = book["Результаты"]
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet.max_column, 10)
        self.assertAlmostEqual(
            sheet.cell(2, 6).value,
            self.experiment.state.points[0].measured_modB_mT,
            places=13,
        )
        self.assertIn(
            "Среднее абсолютное отклонение, мТл",
            [row[0].value for row in book["Параметры"]],
        )
        book.close()
        self.experiment.start_automatic()
        with self.assertRaises(ValueError):
            experiment_data(self.experiment)
        self.assertFalse(self.window.experiment.export_button.isEnabled())

    def test_study_export_uses_current_parameters_and_axis(self):
        self.study.set_parameters(15, 0.1, -0.002)
        for axis in SensorStudyAxis:
            self.study.set_axis(axis)
            data = study_data(self.study)
            self.assertEqual(self.study.curve.axis, axis)
            self.assertGreater(len(data.rows), 1)
            parameters = dict(data.parameters)
            self.assertEqual(parameters["Ток I, мА"], 15)
            self.assertEqual(parameters["Толщина d, мм"], 0.1)
            self.assertEqual(parameters["Коэффициент Холла R_H, м³/Кл"], -0.002)
        self.window.study._cursor(-50, -40)
        self.assertIn("Курсор:", self.window.study.cursor.text())
        self.assertIn("Датчик:", self.window.study.cursor.text())

    def test_nonfinite_data_and_formula_like_names_are_text(self):
        destination = self.directory / "values.xlsx"
        data = ExportData(
            "Проверка",
            ["Значение"],
            [[float("nan")], [float("inf")], [1.23456789012345]],
            [("Тип", "=1+1")],
        )
        save_workbook(destination, data)
        book = load_workbook(destination)
        self.assertEqual(book["Результаты"]["A2"].value, "Недопустимое значение")
        self.assertEqual(book["Параметры"]["B3"].data_type, "s")
        self.assertAlmostEqual(
            book["Результаты"]["A4"].value, 1.23456789012345, places=14
        )
        book.close()

    def test_failed_write_keeps_existing_file(self):
        destination = self.directory / "keep.xlsx"
        destination.write_bytes(b"existing")

        def fail(path):
            path.write_bytes(b"incomplete")
            raise OSError("write failed")

        with self.assertRaises(OSError):
            atomic_write(destination, fail)
        self.assertEqual(destination.read_bytes(), b"existing")
        self.assertEqual(list(self.directory.glob(".hallsim-*")), [])

    def test_png_exports_preserve_selection_and_ranges(self):
        wait_until(lambda: self.window.visualization.canvas.grid is not None)
        self.visual.set_path(-40, 20, 40, 20)
        wait_until(lambda: self.window.visualization.path_plot.profile is not None)
        canvas = self.window.visualization.canvas
        canvas.objects.set_selection(("magnet", 1))
        ranges = canvas.view_box.viewRange()
        for name, prepare in (
            ("field", self.window.visualization._prepare_field_image),
            ("profile", self.window.visualization._prepare_profile_image),
            ("study", self.window.study._prepare_image_export),
        ):
            destination = self.directory / f"{name}.png"
            prepare()(destination)
            image = QImage(str(destination))
            self.assertFalse(image.isNull())
            self.assertGreater(image.width(), 500)
        self.assertEqual(canvas.objects.selection, ("magnet", 1))
        self.assertEqual(canvas.view_box.viewRange(), ranges)
        self.measure()
        self.experiment.compare()
        self.window.experiment.table_toggle.setChecked(True)
        self.window.experiment._prepare_image_export()(
            self.directory / "experiment.png"
        )
        self.assertFalse(QImage(str(self.directory / "experiment.png")).isNull())
        self.assertFalse(self.window.experiment.plot.isVisible())

    def test_hidden_canvas_export_uses_updated_layout_size(self):
        wait_until(lambda: self.window.visualization.canvas.grid is not None)
        visual = self.window.visualization
        previous_height = visual.canvas.height()
        self.window.set_section(1)
        visual.collapse_plot.setChecked(True)
        destination = self.directory / "hidden-map.png"
        visual._prepare_field_image()(destination)
        image = QImage(str(destination))
        self.assertFalse(image.isNull())
        self.assertGreater(image.height(), (previous_height + 48) * 2)
        self.assertTrue(visual.collapse_plot.isChecked())
        self.assertEqual(self.window.pages.currentIndex(), 1)

    def test_settings_recover_invalid_scene_and_preserve_layers(self):
        for payload in (
            [],
            None,
            {"version": 1},
            {**SettingsService._scene_to_dict(single_magnet()), "magnets": []},
        ):
            self.settings._scene_path.write_text(json.dumps(payload))
            loaded = self.settings.load()
            self.assertIsNone(loaded.last_scene)
            self.assertTrue(self.settings.warnings)
        self.visual.set_layer_visibility("field_lines", True)
        self.settings.save(
            self.visual.app_settings(), self.visual.scene(), self.window.saveGeometry()
        )
        loaded = self.settings.load()
        self.assertTrue(loaded.last_scene.view.show_field_lines)
        self.assertTrue(loaded.last_scene.view.show_field_arrows)
        self.settings._settings.setValue("view/zoom", "nan")
        self.assertEqual(self.settings.load().last_zoom, 1)

    def test_manual_overwrite_and_reset_behavior_unchanged(self):
        self.experiment.edit_scene("set_path", -40, 20, 40, 20)
        self.experiment.activate_manual_point(0, -40, 20)
        self.experiment.capture()
        with patch.object(QMessageBox, "question") as question:
            self.window.experiment._reset()
            question.assert_not_called()
        self.assertEqual(self.experiment.captures, [])
        self.experiment.activate_manual_point(0, -40, 20)
        self.experiment.capture()
        self.experiment.finish_manual_point()
        self.assertEqual(self.experiment.state.points[0].source.name, "MANUAL")
        self.experiment.set_config(3, 10)
        self.experiment.start_automatic()
        wait_until(lambda: not self.experiment.running)
        self.assertEqual(self.experiment.state.points[0].source.name, "AUTOMATIC")

    def test_singular_measurement_behavior_unchanged(self):
        self.experiment.edit_scene("set_path", -20, 0, 20, 0)
        self.experiment.set_config(3, 10)
        self.experiment.start_automatic()
        wait_until(lambda: not self.experiment.running)
        self.experiment.compare()
        self.assertEqual(self.experiment.progress, (3, 3))
        self.assertFalse(
            math.isfinite(self.experiment.state.points[1].measured_modB_mT)
        )
        self.assertEqual(self.experiment.comparison.mean_abs_error_mT, 0)


class FieldCalculationTests(unittest.TestCase):
    def test_superseded_drag_preview_is_shown_but_not_after_release(self):
        service = CalculationService()
        previews, final = [], []
        service.fieldPreviewReady.connect(previews.append)
        service.fieldGridReady.connect(final.append)
        started = [Event(), Event()]
        released = [Event(), Event()]
        compute = service.engine.compute_field_grid
        calls = []

        def delayed(*args, **kwargs):
            index = len(calls)
            calls.append(kwargs["scene_revision"])
            if index < 2:
                started[index].set()
                released[index].wait(5)
            return compute(*args, **kwargs)

        with patch.object(service.engine, "compute_field_grid", side_effect=delayed):
            try:
                scene = single_magnet()
                scene.revision = 1
                service.request_field_grid(scene, "drag")
                wait_until(started[0].is_set)
                scene.magnets[0].x_mm = 10
                scene.revision = 2
                service.request_field_grid(scene, "drag")
                released[0].set()
                wait_until(lambda: len(previews) == 1 and started[1].is_set())
                self.assertEqual(previews[0].scene_revision, 1)
                self.assertEqual(previews[0].bx_mT.shape, (200, 320))
                self.assertEqual(final, [])
                scene.magnets[0].x_mm = 20
                scene.revision = 3
                service.request_field_grid(scene, "full")
                released[1].set()
                wait_until(lambda: bool(final))
                self.assertEqual(len(previews), 1)
                self.assertEqual(final[0].scene_revision, 3)
                self.assertEqual(final[0].bx_mT.shape, (400, 640))
                self.assertFalse(final[0].bx_mT.flags.writeable)
            finally:
                for event in released:
                    event.set()
                service.shutdown()
                APP.processEvents()

    def test_lines_do_not_cross_thin_invalid_regions_with_long_steps(self):
        axis = np.linspace(0, 10, 101)
        valid = np.ones((101, 101), dtype=bool)
        valid[:, 50] = False
        horizontal = np.where(valid, 1.0, np.nan)
        vertical = np.where(valid, 0.0, np.nan)
        grid = FieldGrid(axis, axis, horizontal, vertical, horizontal, valid, 0)
        sampler = GridSampler(grid)
        np.testing.assert_array_equal(
            sampler.segment_valid(
                np.array([[4.1, 2.5], [1.1, 2.5], [6.1, 2.5]]),
                np.array([[6.1, 2.5], [3.1, 2.5], [8.1, 2.5]]),
            ),
            [False, True, True],
        )
        lines = StreamlineBuilder().build(grid, step_mm=0.6)
        self.assertTrue(lines)
        for line in lines:
            points = line.points_mm
            self.assertTrue(np.isfinite(points).all())
            self.assertTrue((points[:, 0] < 5).all() or (points[:, 0] > 5).all())
            self.assertTrue(sampler.segment_valid(points[:-1], points[1:]).all())
            np.testing.assert_allclose(points[:, 1], points[0, 1])


if __name__ == "__main__":
    unittest.main()
